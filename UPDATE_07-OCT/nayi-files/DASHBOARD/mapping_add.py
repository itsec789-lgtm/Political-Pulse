# -*- coding: utf-8 -*-
"""
mapping_add.py  --  naye handle ki mapping Mapping.xlsx me jodta hai,
                    phir rebuild_mapping.py --apply chalata hai.

Admin Panel ka "Handle Mapping" tab ise chalata hai (server.js se).

KYUN EXCEL ME, SEEDHA handle_master.json ME KYUN NAHI:
    rebuild_mapping.py ki pehli line hi kehti hai --
      "Mapping.xlsx hi MASTER hai. Purani mapping ka koi data nahi uthaya jaata."
    Yaani sirf handle_master.json me likha hua AGLI REBUILD PAR UD JAATA.
    Isliye asli jagah Excel hai; JSON uske baad rebuild se khud ban jaata hai
    (DASHBOARD aur UP-DASHBOARD dono ke).

CHALAO:
    python mapping_add.py --json rows.json            # jodo + rebuild
    python mapping_add.py --json rows.json --dry-run  # sirf jaanch
    python mapping_add.py --json rows.json --no-rebuild

rows.json:
    {"rows": [
       {"Twitter Account": "someHandle",     # ya poora x.com link
        "State": "Uttar Pradesh", "AC/PC No": "AC_12", "AC Name": "SADAR",
        "PC Name": "-", "Candidate Name": "Ram Kumar", "Designation": "MLA",
        "Category": "MLA", "Party": "BJP",
        "Categories": "MLA, Top Leaders"}    # optional: har extra ke liye ek
     ]}                                      # aur row jud jaati hai

JAWAAB (stdout, JSON):
    {"ok": true, "added": 2, "rows_written": 3, "skipped": [...],
     "errors": [...], "backup": "...", "rebuild": {...}}

SURAKSHA:
  * Excel pehle backup hoti hai (Mapping/_backup/).
  * Likhna pehle temp file me, phir os.replace -- beech me bijli gayi to
    bhi asli Mapping.xlsx saabut rehti hai.
  * Jo handle pehle se Excel me hai usse HAATH NAHI LAGATE (skip + report).
    4,422 handle ki mapping ek galat upload se bigadni nahi chahiye.
"""
import os, sys, json, shutil, subprocess, re
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, '..', 'Mapping', 'Mapping.xlsx')
BAKDIR = os.path.join(HERE, '..', 'Mapping', '_backup')
REBUILD = os.path.join(HERE, 'rebuild_mapping.py')
SHEET = 'Sheet1'

# Nayi Mapping.xlsx ke 10 column. Purani 9-column Excel bhi chalti hai --
# neeche likhte waqt sirf WAHI column bhare jaate hain jo sheet me sach me
# hain (hdr se dekh kar), isliye dono me kaam karta hai.
COLS = ['State', 'AC/PC No', 'AC/PC NAME', 'Candidate Name', 'Gender',
        'Political Party', 'Category', 'Party Designation',
        'Admin Designation', 'Twitter Account']
# purani Excel ke naam -> naye naam (padhte waqt dono chalein)
ALIAS = {'Political Party': 'Party', 'Party Designation': 'Designation',
         'AC/PC NAME': 'AC Name'}
# dash-ok : sirf ye do zaroori. Baaki column me '-' normal hai --
# media house / official page ka na State hota hai, na seat, na party.
REQUIRED = ['Candidate Name', 'Category']
BLANK = ('', '-', '—', 'nan', 'none', 'null', '(blank)')


def clean(v):
    s = str(v if v is not None else '').strip()
    return '' if s.lower() in BLANK else s


def dash(v):
    return clean(v) or '-'


def handle_of(u):
    """rebuild_mapping.handle_of jaisa hi -- URL/@handle se saaf handle."""
    u = str(u or '').strip()
    if not u:
        return ''
    if 'instagram.com' in u.lower() or 'facebook.com' in u.lower():
        return ''
    u = u.rstrip('/')
    if 'x.com/' in u or 'twitter.com/' in u:
        u = u.split('/')[-1]
    u = u.lstrip('@').split('?')[0].strip()
    return u if re.match(r'^[A-Za-z0-9_]{1,15}$', u) else ''


def die(msg, **extra):
    out = {'ok': False, 'error': msg}
    out.update(extra)
    sys.stdout.write(json.dumps(out))
    sys.exit(1)


def main():
    args = sys.argv[1:]
    if '--json' not in args:
        die('--json <file> chahiye')
    src = args[args.index('--json') + 1]
    dry = '--dry-run' in args
    no_rebuild = '--no-rebuild' in args

    try:
        req = json.load(open(src, encoding='utf-8'))
    except Exception as e:
        die('rows file padhi nahi gayi: %s' % e)
    rows = req.get('rows') or []
    if not rows:
        die('ek bhi row nahi di gayi')

    if not os.path.exists(XLSX):
        die('Mapping.xlsx nahi mila: %s' % os.path.abspath(XLSX))

    try:
        import openpyxl
    except ImportError:
        die('openpyxl nahi hai. Chalao:  pip install openpyxl')

    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]

    # ---- header padho, column ka number nikalo ----
    hdr = {}
    for j, c in enumerate(ws[1], start=1):
        name = clean(c.value)
        if name:
            hdr[name] = j
    # Nayi Excel me ye 10, purani me 9 -- isliye dono me se EK poora mile
    # to chal jaata hai. (Pehle sirf ek list dekhta tha aur nayi Excel par
    # "column nahi mila" bol kar ruk jaata tha.)
    OLD_COLS = ['State', 'AC/PC No', 'AC Name', 'PC Name', 'Candidate Name',
                'Designation', 'Category', 'Party', 'Twitter Account']
    missing_cols = [c for c in COLS if c not in hdr]
    if missing_cols:
        miss_old = [c for c in OLD_COLS if c not in hdr]
        if not miss_old:
            missing_cols = []        # purani Excel hai -- wo bhi theek hai
    if missing_cols:
        die('Mapping.xlsx me ye column nahi mile: %s' % ', '.join(missing_cols),
            found=sorted(hdr.keys()))

    # ---- pehle se kaunse handle hain ----
    have = {}
    ucol = hdr['Twitter Account']
    for i in range(2, ws.max_row + 1):
        h = handle_of(ws.cell(row=i, column=ucol).value)
        if h:
            have.setdefault(h.lower(), i)

    # ---- har row jaancho ----
    todo, skipped, errors = [], [], []
    seen = set()
    for n, r in enumerate(rows, start=1):
        h = handle_of(r.get('Twitter Account') or r.get('Handle') or '')
        if not h:
            errors.append({'row': n, 'handle': str(r.get('Twitter Account') or r.get('Handle') or ''),
                           'why': 'handle samajh nahi aaya'})
            continue
        if h.lower() in have:
            skipped.append({'handle': h, 'why': 'Mapping.xlsx me pehle se hai (excel row %d)'
                            % have[h.lower()]})
            continue
        if h.lower() in seen:
            skipped.append({'handle': h, 'why': 'isi file me do baar aaya'})
            continue
        miss = [c for c in REQUIRED if not clean(r.get(c))]
        if miss:
            errors.append({'row': n, 'handle': h, 'why': 'khaali column: ' + ', '.join(miss)})
            continue
        seen.add(h.lower())

        # Categories: har extra category ke liye ek aur Excel row --
        # rebuild_mapping handle ki saari rows ke Category se 'categories'
        # banata hai, isliye multi-category ka yahi sahi tarika hai.
        cats = [clean(r.get('Category'))]
        for x in str(r.get('Categories') or '').split(','):
            x = clean(x)
            if x and x not in cats:
                cats.append(x)
        todo.append({'handle': h, 'row': r, 'cats': cats})

    written = sum(len(t['cats']) for t in todo)
    result = {'ok': True, 'added': len(todo), 'rows_written': written,
              'handles': [t['handle'] for t in todo],
              'skipped': skipped, 'errors': errors, 'dry_run': dry}

    if dry or not todo:
        result['note'] = 'dry-run -- kuch nahi likha' if dry else 'jodne laayak koi row nahi'
        sys.stdout.write(json.dumps(result))
        return

    # ---- backup ----
    os.makedirs(BAKDIR, exist_ok=True)
    bak = os.path.join(BAKDIR, 'Mapping_%s.xlsx' % datetime.now().strftime('%Y%m%d_%H%M%S'))
    shutil.copy2(XLSX, bak)
    result['backup'] = os.path.abspath(bak)

    # ---- rows jodo ----
    at = ws.max_row + 1
    for t in todo:
        r = t['row']
        for cat in t['cats']:
            # request me jo naam aaya ho, ya uska purana roop -- dono se utha lo
            def g(*names):
                for nm in names:
                    v = dash(r.get(nm))
                    if v and v != '-':
                        return v
                return '-'
            seat = g('AC/PC NAME', 'AC Name', 'PC Name')
            vals = {
                'State': g('State'),
                'AC/PC No': g('AC/PC No'),
                'AC/PC NAME': seat,
                'AC Name': seat,          # purani Excel ke liye
                'PC Name': g('PC Name'),
                'Candidate Name': g('Candidate Name'),
                'Gender': g('Gender'),
                'Political Party': g('Political Party', 'Party'),
                'Party': g('Political Party', 'Party'),
                'Category': cat,
                'Party Designation': g('Party Designation', 'Designation'),
                'Admin Designation': g('Admin Designation'),
                'Designation': g('Party Designation', 'Designation'),
                'Twitter Account': 'https://x.com/' + t['handle'],
            }
            # SIRF wahi column likho jo is sheet me sach me hain --
            # isi se nayi (10) aur purani (9) dono Excel chalti hain
            for c, v in vals.items():
                if c in hdr:
                    ws.cell(row=at, column=hdr[c]).value = v
            at += 1

    # ---- temp me likho, phir jagah badlo ----
    tmp = XLSX + '.tmp_%d' % os.getpid()
    try:
        wb.save(tmp)
        os.replace(tmp, XLSX)
    except Exception as e:
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except Exception:
            pass
        die('Excel save nahi hui: %s' % e, backup=result.get('backup'))

    # ---- rebuild: isi se dono handle_master.json bante hain ----
    if not no_rebuild:
        try:
            p = subprocess.run([sys.executable, REBUILD, '--apply'],
                               capture_output=True, text=True, timeout=900,
                               cwd=HERE)
            tail = (p.stdout or '')[-1800:]
            result['rebuild'] = {'ok': p.returncode == 0, 'code': p.returncode,
                                 'out': tail, 'err': (p.stderr or '')[-600:]}
        except Exception as e:
            result['rebuild'] = {'ok': False, 'error': str(e)}
    else:
        result['rebuild'] = {'skipped': True}

    sys.stdout.write(json.dumps(result))


if __name__ == '__main__':
    main()
